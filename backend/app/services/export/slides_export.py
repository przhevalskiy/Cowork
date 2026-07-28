"""
PowerPoint (.pptx) export for CoSight reports — python-pptx based.

CBS-branded slide template:
  - Title slide
  - KPI grid slide (2×3 KPI cards)
  - Narrative/insights slide
  - Data table slide

All slide content comes from the ReportArtifact — never a separate fetch.
Invariant: numbers in the deck ≡ numbers in the artifact.
"""
from __future__ import annotations

import io
from typing import List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from app.models.cosight_report import ReportArtifact, ReportSection

# ── CBS brand palette ─────────────────────────────────────────────────────────
CBS_NAVY = RGBColor(0x00, 0x1B, 0x5E)      # CBS primary navy
CBS_LIGHT_BLUE = RGBColor(0x5B, 0x9B, 0xD5)
CBS_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CBS_DARK_GRAY = RGBColor(0x2D, 0x2D, 0x2D)
CBS_MID_GRAY = RGBColor(0x6B, 0x6B, 0x6B)
CBS_ACCENT = RGBColor(0x00, 0x7A, 0xB3)

# ── Layout constants ──────────────────────────────────────────────────────────
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)


def _add_text_box(slide, left, top, width, height, text, font_size=16,
                  bold=False, color=None, alignment=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color or CBS_DARK_GRAY
    return txBox


def _fill_shape(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _title_slide(prs: Presentation, artifact: ReportArtifact) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Background
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    _fill_shape(bg, CBS_NAVY)

    # Title
    _add_text_box(
        slide, MARGIN, Inches(2.2), Inches(9), Inches(1.2),
        artifact.title, font_size=40, bold=True,
        color=CBS_WHITE, alignment=PP_ALIGN.LEFT,
    )

    # Subtitle / date range
    if artifact.subtitle or artifact.date_range_start:
        sub = artifact.subtitle or f"{artifact.date_range_start} – {artifact.date_range_end}"
        _add_text_box(
            slide, MARGIN, Inches(3.5), Inches(9), Inches(0.7),
            sub, font_size=18, color=CBS_LIGHT_BLUE,
        )

    # CoSight label bottom-right
    _add_text_box(
        slide, Inches(10), Inches(6.8), Inches(3), Inches(0.5),
        "CoSight  |  CBS Marketing & Communications",
        font_size=11, color=CBS_LIGHT_BLUE, alignment=PP_ALIGN.RIGHT,
    )


def _section_header_slide(prs: Presentation, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Dark band
    band = slide.shapes.add_shape(1, 0, Inches(3.0), SLIDE_W, Inches(1.6))
    _fill_shape(band, CBS_NAVY)
    _add_text_box(
        slide, MARGIN, Inches(3.1), Inches(11), Inches(1.2),
        title, font_size=32, bold=True, color=CBS_WHITE, alignment=PP_ALIGN.CENTER,
    )


def _kpi_grid_slide(prs: Presentation, section: ReportSection) -> None:
    """Up to 6 KPI cards in a 2-row × 3-col grid."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Slide title
    _add_text_box(slide, MARGIN, Inches(0.3), Inches(12), Inches(0.6),
                  section.title, font_size=24, bold=True, color=CBS_NAVY)

    kpis = section.kpis[:6]
    cols, rows = 3, 2
    card_w = Inches(3.9)
    card_h = Inches(2.6)
    h_gap = Inches(0.25)
    v_gap = Inches(0.3)
    start_x = MARGIN
    start_y = Inches(1.1)

    for i, kpi in enumerate(kpis):
        col = i % cols
        row = i // cols
        x = start_x + col * (card_w + h_gap)
        y = start_y + row * (card_h + v_gap)

        # Card background
        card = slide.shapes.add_shape(1, x, y, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFB)
        card.line.color.rgb = RGBColor(0xD0, 0xDB, 0xEC)

        label = kpi.label or kpi.metric_key.replace("_", " ").title()
        _add_text_box(slide, x + Inches(0.15), y + Inches(0.15), card_w - Inches(0.3), Inches(0.5),
                      label, font_size=12, color=CBS_MID_GRAY)

        # Value
        val_str = f"{kpi.value:,.0f}" if kpi.unit.value == "count" else f"{kpi.value:.1f}"
        _add_text_box(slide, x + Inches(0.15), y + Inches(0.65), card_w - Inches(0.3), Inches(0.9),
                      val_str, font_size=32, bold=True, color=CBS_NAVY)

        # Delta
        if kpi.delta_pct is not None:
            sign = "+" if kpi.delta_pct > 0 else ""
            delta_color = RGBColor(0x16, 0xA3, 0x4A) if kpi.delta_pct >= 0 else RGBColor(0xDC, 0x26, 0x26)
            _add_text_box(slide, x + Inches(0.15), y + Inches(1.7), card_w - Inches(0.3), Inches(0.5),
                          f"{sign}{kpi.delta_pct:.1f}% vs prior period",
                          font_size=11, color=delta_color)


def _narrative_slide(prs: Presentation, section: ReportSection) -> None:
    if not section.narrative:
        return
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide, MARGIN, Inches(0.3), Inches(12), Inches(0.6),
                  section.title, font_size=24, bold=True, color=CBS_NAVY)
    _add_text_box(slide, MARGIN, Inches(1.1), Inches(12), Inches(5.5),
                  section.narrative, font_size=14, color=CBS_DARK_GRAY, wrap=True)


def _table_slide(prs: Presentation, section: ReportSection) -> None:
    """Render section.table_rows as a simple table slide (max 15 rows)."""
    if not section.table_rows or not section.table_headers:
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text_box(slide, MARGIN, Inches(0.3), Inches(12), Inches(0.6),
                  section.title, font_size=24, bold=True, color=CBS_NAVY)

    headers = section.table_headers
    rows_data = section.table_rows[:15]
    n_cols = len(headers)
    n_rows = len(rows_data) + 1  # +1 for header row

    tbl = slide.shapes.add_table(
        n_rows, n_cols,
        MARGIN, Inches(1.1),
        SLIDE_W - 2 * MARGIN, Inches(5.8),
    ).table

    col_width = (SLIDE_W - 2 * MARGIN) / n_cols

    # Header row
    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CBS_NAVY
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = CBS_WHITE

    # Data rows
    for ri, row_dict in enumerate(rows_data, start=1):
        for ci, key in enumerate(headers):
            cell = tbl.cell(ri, ci)
            cell.text = str(row_dict.get(key, ""))
            if ri % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFB)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = CBS_DARK_GRAY


def export_report_to_pptx(artifact: ReportArtifact) -> bytes:
    """
    Render a ReportArtifact into a CBS-branded .pptx and return raw bytes.
    Slide order: title → (per section: header → KPI grid → narrative → table)
    """
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _title_slide(prs, artifact)

    for section in artifact.sections:
        _section_header_slide(prs, section.title)
        if section.kpis:
            _kpi_grid_slide(prs, section)
        if section.narrative:
            _narrative_slide(prs, section)
        if section.table_rows:
            _table_slide(prs, section)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
